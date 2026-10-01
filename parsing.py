"""
Extracts and reconstructs text from uploaded resume / JD files.

Supports PDF, DOCX, and plain text.
No external service calls and no OCR.

PDF extraction is layout-aware:
- extracts words with coordinates
- reconstructs lines
- detects likely multi-column pages
- preserves reading order as much as possible
"""

import io
import re
from collections import Counter

import pdfplumber
import docx


# ---------------------------------------------------------------------------
# PDF layout helpers
# ---------------------------------------------------------------------------

def _group_words_into_lines(words, y_tolerance=3, page_width=None):
    """Group PDF words into visually aligned text lines."""
    if not words:
        return []

    words = sorted(words, key=lambda w: (w["top"], w["x0"]))
    rows = []

    for word in words:
        if not rows:
            rows.append([word])
            continue

        previous_top = sum(w["top"] for w in rows[-1]) / len(rows[-1])

        if abs(word["top"] - previous_top) <= y_tolerance:
            rows[-1].append(word)
        else:
            rows.append([word])

    # A big horizontal jump between two words that share a row usually means
    # they come from different columns of a multi-column layout (e.g. two
    # side-by-side section headers like "WORK EXPERIENCE" and "EDUCATION"),
    # not the same line of text. Split on such gaps so they don't get fused
    # into one line that matches no known section header, which silently
    # breaks section detection downstream. Words are re-sorted by x0 within
    # each row first, since the initial (top, x0) sort can otherwise place a
    # right-column word ahead of a left-column word on the same row.
    gap_threshold = (page_width * 0.15) if page_width else 60
    lines = []

    for row_words in rows:
        row_words = sorted(row_words, key=lambda w: w["x0"])
        current = [row_words[0]]

        for word in row_words[1:]:
            previous_right = max(w["x1"] for w in current)

            if word["x0"] - previous_right > gap_threshold:
                lines.append(current)
                current = [word]
            else:
                current.append(word)

        lines.append(current)

    result = []

    for line_words in lines:
        line_words.sort(key=lambda w: w["x0"])
        text = " ".join(w["text"] for w in line_words).strip()

        if text:
            result.append({
                "text": text,
                "x0": min(w["x0"] for w in line_words),
                "x1": max(w["x1"] for w in line_words),
                "top": min(w["top"] for w in line_words),
                "bottom": max(w["bottom"] for w in line_words),
            })

    return result


def _detect_column_split(lines, page_width):
    """
    Detect a strong vertical gap suggesting a two-column page.

    Returns the x-coordinate of the split, or None when the page appears
    to be single-column.
    """
    if len(lines) < 8:
        return None

    gaps = []

    for line in lines:
        if line["x1"] < page_width:
            gaps.append((line["x0"], line["x1"]))

    if not gaps:
        return None

    candidate_gaps = []

    for left in lines:
        for right in lines:
            if right["x0"] <= left["x1"]:
                continue

            gap = right["x0"] - left["x1"]

            if gap >= page_width * 0.08:
                candidate_gaps.append(
                    (gap, (left["x1"] + right["x0"]) / 2)
                )

    if not candidate_gaps:
        return None

    largest_gap, split = max(candidate_gaps, key=lambda x: x[0])

    left_count = sum(line["x1"] <= split for line in lines)
    right_count = sum(line["x0"] >= split for line in lines)

    # Require meaningful content on both sides.
    if left_count >= 4 and right_count >= 4:
        return split

    return None


def _reconstruct_page(page):
    """
    Reconstruct a PDF page in reading order.

    Single-column pages are returned top-to-bottom.
    Strong two-column pages are returned left-column then right-column.
    """
    words = page.extract_words(
        x_tolerance=2,
        y_tolerance=3,
        keep_blank_chars=False,
        use_text_flow=False,
    )

    if not words:
        return ""

    lines = _group_words_into_lines(words, page_width=page.width)

    split = _detect_column_split(lines, page.width)

    if split is None:
        lines.sort(key=lambda line: (line["top"], line["x0"]))
        return "\n".join(line["text"] for line in lines)

    left_lines = [
        line for line in lines
        if line["x1"] <= split
    ]

    right_lines = [
        line for line in lines
        if line["x0"] >= split
    ]

    # If the supposed split leaves too much content in the middle,
    # fall back to ordinary reading order.
    assigned = len(left_lines) + len(right_lines)

    if assigned < len(lines) * 0.85:
        lines.sort(key=lambda line: (line["top"], line["x0"]))
        return "\n".join(line["text"] for line in lines)

    left_lines.sort(key=lambda line: (line["top"], line["x0"]))
    right_lines.sort(key=lambda line: (line["top"], line["x0"]))

    ordered = left_lines + right_lines

    return "\n".join(line["text"] for line in ordered)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = _reconstruct_page(page)

            if page_text.strip():
                text_parts.append(page_text)

    return "\n".join(text_parts)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))

    paragraphs = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Resume templates frequently store education/experience in tables.
    for table in doc.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                paragraphs.append(" | ".join(cells))

    return "\n".join(paragraphs)


# ---------------------------------------------------------------------------
# Generic dispatcher
# ---------------------------------------------------------------------------

def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Dispatch based on file extension.

    Returns reconstructed text or an empty string if parsing fails.
    """
    lower = filename.lower()

    try:
        if lower.endswith(".pdf"):
            return extract_text_from_pdf(file_bytes)

        elif lower.endswith(".docx"):
            return extract_text_from_docx(file_bytes)

        elif lower.endswith(".txt"):
            return file_bytes.decode("utf-8", errors="ignore")

        else:
            return ""

    except Exception:
        return ""